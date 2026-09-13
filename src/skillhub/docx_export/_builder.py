"""Собирает Word-документ только из текстовых узлов модели протокола."""

from collections.abc import Sequence
from io import BytesIO
from typing import Protocol as TypingProtocol

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table

from skillhub.docx_export._placeholders import (
    EMPTY_SECTION,
    assignee_label,
    due_label,
    participants_label,
    section_items,
)
from skillhub.docx_export._sanitize import sanitize_docx_text
from skillhub.protocol.models import Protocol


class _TextCell(TypingProtocol):
    """Описывает ячейку таблицы с текстовым узлом."""

    text: str


_TITLE_PREFIX = "Протокол встречи:"
_DATE_LABEL = "Дата:"
_PARTICIPANTS_LABEL = "Участники:"
_DISCUSSION = "Обсуждение"
_DECISIONS = "Решения"
_TASKS = "Задачи"
_QUESTIONS = "Открытые вопросы"
_TASK_HEADERS = ("Задача", "Ответственный", "Срок")


def build_docx(protocol: Protocol) -> bytes:
    """Собирает байты Word-документа из нормативной модели протокола.

    Args:
        protocol: принятая модель протокола без сырого Markdown.

    Returns:
        Байтовое представление безопасного документа Word.
    """
    document = Document()
    _add_header(document, protocol)
    _add_item_section(document, _DISCUSSION, protocol.discussion)
    _add_item_section(document, _DECISIONS, protocol.decisions)
    _add_tasks_section(document, protocol)
    _add_item_section(document, _QUESTIONS, protocol.open_questions)
    return _document_bytes(document)


def _add_header(document: DocumentObject, protocol: Protocol) -> None:
    title = sanitize_docx_text(f"{_TITLE_PREFIX} {protocol.title}")
    document.add_heading(title, level=1)
    document.add_paragraph(sanitize_docx_text(f"{_DATE_LABEL} {protocol.date}"))
    participants = participants_label(protocol.participants)
    document.add_paragraph(
        sanitize_docx_text(f"{_PARTICIPANTS_LABEL} {participants}"),
    )


def _add_item_section(
    document: DocumentObject,
    title: str,
    items: tuple[str, ...],
) -> None:
    document.add_heading(sanitize_docx_text(title), level=2)
    _add_item_body(document, section_items(items))


def _add_item_body(document: DocumentObject, items: tuple[str, ...]) -> None:
    for item in items:
        document.add_paragraph(sanitize_docx_text(item), style="List Bullet")


def _add_tasks_section(document: DocumentObject, protocol: Protocol) -> None:
    document.add_heading(sanitize_docx_text(_TASKS), level=2)
    if not protocol.tasks:
        document.add_paragraph(sanitize_docx_text(EMPTY_SECTION), style="List Bullet")
        return
    table = document.add_table(rows=1 + len(protocol.tasks), cols=3)
    table.style = "Table Grid"
    _fill_tasks_table(table, protocol)
    _apply_visible_borders(table)


def _fill_tasks_table(table: Table, protocol: Protocol) -> None:
    _write_row(table.rows[0].cells, _TASK_HEADERS)
    for index, task in enumerate(protocol.tasks, start=1):
        _write_row(
            table.rows[index].cells,
            (
                task.title,
                assignee_label(task.assignee),
                due_label(task.due),
            ),
        )


def _write_row(cells: Sequence[_TextCell], values: tuple[str, ...]) -> None:
    for cell, value in zip(cells, values, strict=True):
        cell.text = sanitize_docx_text(value)


def _apply_visible_borders(table: Table) -> None:
    tbl_pr = table._tbl.tblPr  # noqa: SLF001
    existing = tbl_pr.find(qn("w:tblBorders"))
    if existing is not None:
        tbl_pr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        line = OxmlElement(f"w:{edge}")
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), "8")
        line.set(qn("w:space"), "0")
        line.set(qn("w:color"), "000000")
        borders.append(line)
    tbl_pr.append(borders)


def _document_bytes(document: DocumentObject) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()

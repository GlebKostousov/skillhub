# ruff: noqa: RUF001
"""Проверяет HTTP-швы страницы протокола и выгрузки Word."""

from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from fastapi.testclient import TestClient
from httpx2 import Response
from structlog.testing import capture_logs

from skillhub.app_factory import create_app
from skillhub.protocol import GeneratedDraft, parse
from skillhub.web import create_protocol_router, install_error_handlers
from skillhub.web.protocol_routes import MAX_PROTOCOL_BODY_BYTES

_WEB_ROOT = Path(__file__).resolve().parents[1] / "src" / "skillhub" / "web"
_TEMPLATES = _WEB_ROOT / "templates"
_PROTOCOL_JS = _WEB_ROOT / "static" / "protocol.js"
_CSRF = "protocol-csrf-token-value"
_ORIGIN = {"origin": "http://testserver"}
_VALID_MARKDOWN = """# Протокол встречи: Тема

**Дата:** 2026-09-12
**Участники:** Анна

## Обсуждение
- пункт

## Решения
- решение

## Задачи
- —

## Открытые вопросы
- вопрос
"""
_TRAVERSAL_MARKDOWN = """# Протокол встречи: ../../evil.docx

**Дата:** 2026-09-12
**Участники:** Анна

## Обсуждение
- пункт

## Решения
- решение

## Задачи
- —

## Открытые вопросы
- вопрос
"""
_INVALID_MARKDOWN = "это не протокол"


class _CsrfParser(HTMLParser):
    """Собирает скрытые CSRF-токены страницы протокола."""

    def __init__(self) -> None:
        """Создаёт пустой список найденных токенов."""
        super().__init__()
        self.tokens: list[str] = []

    def handle_starttag(
        self,
        tag: str,
        attrs: list[tuple[str, str | None]],
    ) -> None:
        """Запоминает значение скрытого поля CSRF.

        Args:
            tag: имя открывающего HTML-тега.
            attrs: атрибуты открывающего HTML-тега.
        """
        attributes = dict(attrs)
        if tag == "input" and attributes.get("name") == "csrf_token":
            token = attributes.get("value")
            if token:
                self.tokens.append(token)


class _FakeGenerator:
    """Возвращает заранее заданный черновик и запоминает вызовы порта."""

    def __init__(self, draft: GeneratedDraft) -> None:
        """Сохраняет ответ порта.

        Args:
            draft: заранее собранный текстовый ответ.
        """
        self.draft = draft
        self.calls: list[tuple[str, str]] = []

    def generate(self, instruction: str, material: str) -> GeneratedDraft:
        """Возвращает заготовленный черновик.

        Args:
            instruction: доверенная инструкция.
            material: транскрипция встречи.
        """
        self.calls.append((instruction, material))
        return self.draft


def test_protocol_page_renders_draft_and_download_controls() -> None:
    """Проверяет панель протокола на главной и кнопку выгрузки."""
    client = TestClient(create_app())
    missing_page = client.get("/protocol", follow_redirects=False)
    response = client.get("/")
    parser = _CsrfParser()
    parser.feed(response.text)

    assert missing_page.status_code == 404
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/html")
    assert "<h1" in response.text
    assert "Вопросы для полноты данных" in response.text
    assert "Сохранить в Word" in response.text
    assert "Генерация черновика недоступна." not in response.text
    assert 'id="protocol-followup"' in response.text
    assert 'id="protocol-draft-button" disabled' in response.text
    assert 'id="protocol-download-button"' in response.text
    assert 'id="protocol-markdown"' in response.text
    assert "alert-danger" in response.text
    assert 'id="protocol-status"' in response.text
    assert 'id="protocol-status" class="visually-hidden"' not in response.text
    assert "/static/protocol.js" in response.text
    assert len(parser.tokens) == 1
    script = _PROTOCOL_JS.read_text(encoding="utf-8")
    assert "innerHTML" not in script
    assert "textContent" in script
    assert "try" in script
    assert "catch" in script
    assert "Не получилось связаться. Проверьте сеть и попробуйте ещё раз." in script
    assert "Собирается черновик…" in script
    assert "Собираю файл Word…" in script
    assert "error.expected" in script
    assert "error.got" in script
    assert "error.line" in script


def test_protocol_page_enables_draft_when_generation_available() -> None:
    """Проверяет рабочий черновик API при подключённом порте генерации."""
    client = _router_client(
        _FakeGenerator(GeneratedDraft(text=_VALID_MARKDOWN, finish_reason="stop")),
    )
    response = _post_json(
        client,
        "/protocol/draft",
        _CSRF,
        {"transcript": "Анна: обсудили тему."},
    )

    assert response.status_code == 200
    assert "text" in response.json()


def test_draft_without_generator_is_unavailable() -> None:
    """Проверяет отказ черновика, если порт генерации не подключён."""
    client = TestClient(create_app())
    response = _post_json(
        client,
        "/protocol/draft",
        _csrf(client),
        {"transcript": "Анна: обсудили тему."},
    )

    assert response.status_code == 503
    assert response.json() == {
        "error": {
            "code": "generation_unavailable",
            "message": "Генерация черновика недоступна.",
        }
    }
    assert "content-disposition" not in response.headers


def test_draft_returns_protocol_from_generator() -> None:
    """Проверяет сборку черновика через порт генерации без записи на диск."""
    generator = _FakeGenerator(
        GeneratedDraft(text=_VALID_MARKDOWN, finish_reason="stop"),
    )
    client = _router_client(generator)
    transcript = "Анна: согласуем тему встречи."

    response = _post_json(
        client,
        "/protocol/draft",
        _CSRF,
        {"transcript": transcript},
    )

    assert response.status_code == 200
    payload = response.json()
    text = payload["text"]
    assert payload["title"] == "Тема"
    assert isinstance(text, str)
    assert parse(text) == parse(_VALID_MARKDOWN)
    assert generator.calls[0][1] == transcript
    assert generator.calls[0][0] != transcript


def test_docx_download_uses_fixed_filename() -> None:
    """Проверяет выгрузку Word с фиксированным именем файла."""
    client = TestClient(create_app())
    response = _post_json(
        client,
        "/protocol/docx",
        _csrf(client),
        {"text": _TRAVERSAL_MARKDOWN},
    )
    document = Document(BytesIO(response.content))

    assert response.status_code == 200
    disposition = response.headers["content-disposition"]
    assert disposition == 'attachment; filename="protocol.docx"'
    assert "evil" not in response.headers["content-disposition"]
    assert document.paragraphs[0].text == "Протокол встречи: ../../evil.docx"


def test_invalid_markdown_returns_parse_details_without_file() -> None:
    """Проверяет 422 с координатами разбора и отсутствие файла."""
    client = TestClient(create_app())
    with capture_logs() as logs:
        response = _post_json(
            client,
            "/protocol/docx",
            _csrf(client),
            {"text": _INVALID_MARKDOWN},
        )

    error = response.json()["error"]
    assert response.status_code == 422
    assert set(error) == {"code", "message", "line", "expected", "got"}
    assert error["code"] == "protocol_parse_error"
    assert error["message"] == "Структура протокола нарушена."
    assert isinstance(error["line"], int)
    assert isinstance(error["expected"], str)
    assert error["got"] == _INVALID_MARKDOWN
    assert "content-disposition" not in response.headers
    assert _INVALID_MARKDOWN not in repr(logs)
    assert all("expected" not in entry for entry in logs)
    assert all("got" not in entry for entry in logs)


def test_unicode_protocol_survives_http_export() -> None:
    """Проверяет сохранение Unicode при выгрузке через HTTP."""
    markdown = _VALID_MARKDOWN.replace("пункт", "Тема 漢字 и кириллица 🎉")
    client = TestClient(create_app())
    response = _post_json(
        client,
        "/protocol/docx",
        _csrf(client),
        {"text": markdown},
    )
    document = Document(BytesIO(response.content))
    texts = [paragraph.text for paragraph in document.paragraphs]

    assert "Тема 漢字 и кириллица 🎉" in texts


def test_foreign_origin_is_forbidden() -> None:
    """Проверяет отказ изменяющего запроса с чужим Origin."""
    client = TestClient(create_app())
    response = client.post(
        "/protocol/docx",
        json={"csrf_token": _csrf(client), "text": _VALID_MARKDOWN},
        headers={"origin": "https://attacker.example"},
    )

    assert response.status_code == 403
    assert response.json()["error"]["code"] == "protocol_forbidden"
    assert "content-disposition" not in response.headers


def test_wrong_csrf_is_forbidden() -> None:
    """Проверяет отказ при неверном секрете CSRF."""
    client = TestClient(create_app())
    response = _post_json(
        client,
        "/protocol/docx",
        "wrong-csrf-token-value",
        {"text": _VALID_MARKDOWN},
    )

    assert response.status_code == 403
    assert response.json()["error"]["code"] == "protocol_forbidden"


def test_oversized_body_is_rejected() -> None:
    """Проверяет отказ по заявленному размеру тела до разбора."""
    client = TestClient(create_app())
    response = client.post(
        "/protocol/docx",
        content=b"{}",
        headers={
            "origin": "http://testserver",
            "content-type": "application/json",
            "content-length": str(MAX_PROTOCOL_BODY_BYTES + 1),
        },
    )

    assert response.status_code == 413
    assert response.json()["error"]["code"] == "request_too_large"


def test_missing_origin_is_forbidden() -> None:
    """Проверяет отказ изменяющего запроса без заголовка Origin."""
    client = TestClient(create_app())
    response = client.post(
        "/protocol/docx",
        json={"csrf_token": _csrf(client), "text": _VALID_MARKDOWN},
    )

    assert response.status_code == 403
    assert response.json()["error"]["code"] == "protocol_forbidden"


def test_invalid_json_and_missing_text_are_rejected() -> None:
    """Проверяет отказ для битого JSON и отсутствующего текста протокола."""
    client = TestClient(create_app())
    token = _csrf(client)
    broken = client.post(
        "/protocol/docx",
        content=b"{",
        headers={
            "origin": "http://testserver",
            "content-type": "application/json",
        },
    )
    missing = _post_json(client, "/protocol/docx", token, {})

    assert broken.status_code == 422
    assert broken.json()["error"]["code"] == "protocol_invalid_request"
    assert missing.status_code == 422
    assert missing.json()["error"]["code"] == "protocol_invalid_request"


def test_non_json_body_is_rejected() -> None:
    """Проверяет отказ для тела, которое не является JSON."""
    client = TestClient(create_app())
    response = client.post(
        "/protocol/docx",
        content=b"csrf_token=token",
        headers={
            "origin": "http://testserver",
            "content-type": "application/x-www-form-urlencoded",
        },
    )

    assert response.status_code == 415
    assert "content-disposition" not in response.headers


def test_transcript_is_not_written_to_disk(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Проверяет, что транскрипция не сохраняется в рабочем каталоге."""
    monkeypatch.chdir(tmp_path)
    marker = "unique-transcript-disk-guard-e3s4"
    client = TestClient(create_app())
    _post_json(client, "/protocol/draft", _csrf(client), {"transcript": marker})
    leftovers = [
        path
        for path in tmp_path.rglob("*")
        if path.is_file() and marker.encode() in path.read_bytes()
    ]

    assert leftovers == []


def test_draft_parse_error_keeps_extended_envelope() -> None:
    """Проверяет расширенную оболочку разбора на шве черновика."""
    generator = _FakeGenerator(
        GeneratedDraft(text=_INVALID_MARKDOWN, finish_reason="stop"),
    )
    client = _router_client(generator)
    with capture_logs() as logs:
        response = _post_json(
            client,
            "/protocol/draft",
            _CSRF,
            {"transcript": "секретная транскрипция"},
        )

    error = response.json()["error"]
    assert response.status_code == 422
    assert error["code"] == "protocol_parse_error"
    assert error["got"] == _INVALID_MARKDOWN
    assert "секретная транскрипция" not in repr(logs)
    assert "секретная транскрипция" not in response.text


def test_protocol_router_rejects_invalid_configured_csrf() -> None:
    """Проверяет отказ сборки маршрутизатора при неверном секрете."""
    templates = Jinja2Templates(directory=_TEMPLATES)

    with pytest.raises(ValueError, match="CSRF"):
        create_protocol_router(templates, csrf_token="я", generator=None)  # noqa: S106


def _csrf(client: TestClient) -> str:
    """Читает CSRF-токен со страницы протокола.

    Args:
        client: клиент одного экземпляра приложения.
    """
    parser = _CsrfParser()
    parser.feed(client.get("/").text)
    assert len(parser.tokens) == 1
    return parser.tokens[0]


def _post_json(
    client: TestClient,
    path: str,
    token: str,
    fields: dict[str, str],
) -> Response:
    """Отправляет JSON-запрос протокола с Origin и секретом.

    Args:
        client: HTTP-клиент проверяемого приложения.
        path: путь изменяющего маршрута.
        token: ожидаемый секрет CSRF.
        fields: прикладные поля тела запроса.
    """
    return client.post(
        path,
        json={"csrf_token": token, **fields},
        headers=_ORIGIN,
    )


def _router_client(generator: _FakeGenerator) -> TestClient:
    """Собирает изолированный маршрутизатор протокола.

    Args:
        generator: подставной порт генерации.
    """
    templates = Jinja2Templates(directory=_TEMPLATES)
    app = FastAPI()
    app.mount("/static", StaticFiles(directory=_WEB_ROOT / "static"), name="static")
    app.include_router(
        create_protocol_router(
            templates,
            csrf_token=_CSRF,
            generator=generator,
        )
    )
    install_error_handlers(app)
    return TestClient(app)

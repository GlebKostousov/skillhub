"""Проверяет сборку Word-документа из нормативной модели протокола."""

import ast
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from hypothesis import given, settings
from hypothesis import strategies as st

from skillhub.docx_export import build_docx
from skillhub.docx_export._placeholders import (
    EMPTY_SECTION,
    UNASSIGNED,
    UNSPECIFIED_DUE,
)
from skillhub.protocol import PLACEHOLDER, Protocol, ProtocolTask

_LONG_ITEM = ("Тема " * 400) + "漢字 и кириллица 🎉"
_CONTROLLED_TITLE = "Тема\x00\x08\x1b без управления"


def test_example_protocol_becomes_openable_docx_structure() -> None:
    """Проверяет заголовки, таблицу задач и тексты ячеек через python-docx."""
    protocol = _example_protocol()
    document = _open(build_docx(protocol))

    assert _heading_texts(document) == [
        "Протокол встречи: Подготовка квартального отчёта",
        "Обсуждение",
        "Решения",
        "Задачи",
        "Открытые вопросы",
    ]
    assert "Дата: 2026-09-12" in _paragraph_texts(document)
    assert "Участники: Анна, Борис" in _paragraph_texts(document)
    assert document.tables[0].rows[0].cells[0].text == "Задача"
    assert document.tables[0].rows[1].cells[0].text == "Проверить показатели продаж"
    assert document.tables[0].rows[1].cells[1].text == "Анна"
    assert document.tables[0].rows[2].cells[1].text == UNASSIGNED
    assert document.tables[0].rows[2].cells[2].text == "до следующей встречи"


def test_empty_sections_use_word_placeholders() -> None:
    """Проверяет заполнители Word для пустых разделов и полей задачи."""
    protocol = Protocol(
        title="Тема",
        date=PLACEHOLDER,
        participants=(),
        discussion=(),
        decisions=(),
        tasks=(
            ProtocolTask(title="Сделать отчёт", assignee=PLACEHOLDER, due=PLACEHOLDER),
        ),
        open_questions=(),
    )
    document = _open(build_docx(protocol))
    texts = _paragraph_texts(document)

    assert texts.count(EMPTY_SECTION) == 3
    assert document.tables[0].rows[1].cells[1].text == UNASSIGNED
    assert document.tables[0].rows[1].cells[2].text == UNSPECIFIED_DUE
    assert "Участники: —" in texts


def test_empty_tasks_section_has_placeholder_and_no_table() -> None:
    """Проверяет пустой раздел задач без таблицы Word."""
    document = _open(build_docx(_empty_protocol()))

    assert document.tables == []
    assert EMPTY_SECTION in _paragraph_texts(document)


def test_unicode_and_long_text_are_preserved() -> None:
    """Проверяет сохранение длинного Unicode-текста в текстовых узлах."""
    protocol = _empty_protocol().model_copy(update={"discussion": (_LONG_ITEM,)})
    document = _open(build_docx(protocol))

    assert _LONG_ITEM in _paragraph_texts(document)


def test_control_characters_are_removed_from_text_nodes() -> None:
    """Проверяет удаление запрещённых управляющих символов XML."""
    protocol = _empty_protocol().model_copy(update={"title": _CONTROLLED_TITLE})
    document = _open(build_docx(protocol))

    heading = _heading_texts(document)[0]
    assert "\x00" not in heading
    assert "\x08" not in heading
    assert "\x1b" not in heading
    assert "Тема без управления" in heading


def test_traversal_in_title_stays_plain_text() -> None:
    """Проверяет, что путь в теме остаётся текстом заголовка."""
    protocol = _empty_protocol().model_copy(update={"title": "../../evil.docx"})
    document = _open(build_docx(protocol))

    assert _heading_texts(document)[0] == "Протокол встречи: ../../evil.docx"


def test_repeated_export_keeps_logical_structure() -> None:
    """Проверяет повторный экспорт той же модели без смены структуры."""
    protocol = _example_protocol()

    assert _logical_structure(build_docx(protocol)) == _logical_structure(
        build_docx(protocol)
    )


def test_document_has_only_internal_text_relationships() -> None:
    """Проверяет отсутствие внешних ссылок и гиперссылок в документе."""
    document = _open(build_docx(_example_protocol()))
    hyperlinks = [
        hyperlink.address
        for paragraph in document.paragraphs
        for hyperlink in paragraph.hyperlinks
    ]
    external = [
        rel.target_ref for rel in document.part.rels.values() if rel.is_external
    ]

    assert hyperlinks == []
    assert external == []


def test_docx_export_imports_only_public_protocol_model() -> None:
    """Проверяет, что пакет экспорта берёт модель без parser и renderer."""
    imported = [
        module
        for source_path in _docx_sources()
        for module in _imported_modules(source_path)
    ]
    violations = [
        f"{source_path.name}:{module}"
        for source_path in _docx_sources()
        for module in _imported_modules(source_path)
        if _is_forbidden_docx_import(module)
    ]

    assert "skillhub.protocol.models" in imported
    assert violations == []


_SAFE_CHARS = st.characters(
    whitelist_categories=("L", "N"),
    whitelist_characters=" -.",
    blacklist_characters="|#\n\r,",
)
_TOKEN = (
    st.text(_SAFE_CHARS, min_size=1, max_size=24)
    .map(str.strip)
    .filter(lambda token: token not in {"", PLACEHOLDER})
)
_TASK = st.builds(ProtocolTask, title=_TOKEN, assignee=_TOKEN, due=_TOKEN)
_PROTOCOLS = st.builds(
    Protocol,
    title=_TOKEN,
    date=st.one_of(st.just(PLACEHOLDER), _TOKEN),
    participants=st.lists(_TOKEN, max_size=4).map(tuple),
    discussion=st.lists(_TOKEN, max_size=4).map(tuple),
    decisions=st.lists(_TOKEN, max_size=4).map(tuple),
    tasks=st.lists(_TASK, max_size=3).map(tuple),
    open_questions=st.lists(_TOKEN, max_size=4).map(tuple),
)


@given(_PROTOCOLS)
@settings(max_examples=20)
def test_replay_export_matches_logical_structure(protocol: Protocol) -> None:
    """Проверяет совпадение логической структуры при повторном экспорте."""
    assert _logical_structure(build_docx(protocol)) == _logical_structure(
        build_docx(protocol)
    )


def _example_protocol() -> Protocol:
    """Собирает нормативный пример протокола с задачей-заполнителем."""
    return Protocol(
        title="Подготовка квартального отчёта",
        date="2026-09-12",
        participants=("Анна", "Борис"),
        discussion=(
            "Команда сверила состав квартального отчёта и доступность исходных данных.",
        ),
        decisions=("Включить в отчёт показатели продаж и поддержки.",),
        tasks=(
            ProtocolTask(
                title="Проверить показатели продаж",
                assignee="Анна",
                due="2026-09-15",
            ),
            ProtocolTask(
                title="Подготовить данные поддержки",
                assignee=PLACEHOLDER,
                due="до следующей встречи",
            ),
        ),
        open_questions=("Нужно определить ответственного за данные поддержки.",),
    )


def _empty_protocol() -> Protocol:
    """Собирает протокол с пустыми разделами."""
    return Protocol(
        title="Тема",
        date=PLACEHOLDER,
        participants=(),
        discussion=(),
        decisions=(),
        tasks=(),
        open_questions=(),
    )


def _open(payload: bytes) -> DocumentObject:
    """Открывает собранные байты как документ python-docx.

    Args:
        payload: результат ``build_docx``.

    Returns:
        Документ для проверки структуры.
    """
    return Document(BytesIO(payload))


def _heading_texts(document: DocumentObject) -> list[str]:
    """Возвращает тексты заголовков документа.

    Args:
        document: открытый Word-документ.
    """
    texts: list[str] = []
    for paragraph in document.paragraphs:
        style = paragraph.style
        if style is not None and style.name.startswith("Heading"):
            texts.append(paragraph.text)
    return texts


def _paragraph_texts(document: DocumentObject) -> list[str]:
    """Возвращает тексты всех абзацев документа.

    Args:
        document: открытый Word-документ.
    """
    return [paragraph.text for paragraph in document.paragraphs]


def _logical_structure(
    payload: bytes,
) -> tuple[tuple[str, ...], tuple[tuple[tuple[str, ...], ...], ...]]:
    """Снимает логическую структуру документа для сравнения повторов.

    Args:
        payload: результат ``build_docx``.
    """
    document = _open(payload)
    paragraphs = tuple(_paragraph_texts(document))
    tables = tuple(
        tuple(tuple(cell.text for cell in row.cells) for row in table.rows)
        for table in document.tables
    )
    return paragraphs, tables


def _docx_sources() -> tuple[Path, ...]:
    """Собирает исходники пакета экспорта."""
    root = Path(__file__).resolve().parents[1] / "src" / "skillhub" / "docx_export"
    return tuple(root.rglob("*.py"))


def _imported_modules(source_path: Path) -> tuple[str, ...]:
    """Возвращает имена импортированных модулей файла.

    Args:
        source_path: путь исходника Python.
    """
    tree = ast.parse(source_path.read_text(encoding="utf-8"))
    return tuple(name for node in ast.walk(tree) for name in _import_names(node))


def _import_names(node: ast.AST) -> tuple[str, ...]:
    """Возвращает имена модулей одного узла импорта.

    Args:
        node: узел синтаксического дерева.
    """
    if isinstance(node, ast.Import):
        return tuple(alias.name for alias in node.names)
    if isinstance(node, ast.ImportFrom):
        return (node.module or "",)
    return ()


def _is_forbidden_docx_import(module: str) -> bool:
    """Проверяет запрещённый импорт пакета экспорта.

    Args:
        module: полное имя импортированного модуля.
    """
    forbidden = {
        "skillhub.protocol",
        "skillhub.protocol._parser",
        "skillhub.protocol._draft",
        "skillhub.protocol._renderer",
        "skillhub.protocol._models",
        "skillhub.web",
        "skillhub.registry",
        "skillhub.llm",
    }
    return module in forbidden or module.startswith("skillhub.web.")
